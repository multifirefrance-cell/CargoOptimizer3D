"""Pruebas de `templates.py`: las cuatro plantillas oficiales se generan y se auto-importan."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from cargo_optimizer.infrastructure.excel.catalog_importer import import_catalog
from cargo_optimizer.infrastructure.excel.loading_space_importer import import_loading_spaces
from cargo_optimizer.infrastructure.excel.packing_list_importer import import_packing_list
from cargo_optimizer.infrastructure.excel.templates import (
    CATALOG_TEMPLATE_FILENAME,
    LOADING_SPACE_TEMPLATE_FILENAME,
    OPTIMIZATION_RESULT_TEMPLATE_FILENAME,
    PACKING_LIST_TEMPLATE_FILENAME,
    generate_all_templates,
)


def test_generate_all_templates_creates_the_four_official_files(tmp_path: Path) -> None:
    paths = generate_all_templates(tmp_path, application_version="9.9.9")

    names = {path.name for path in paths}
    assert names == {
        CATALOG_TEMPLATE_FILENAME,
        PACKING_LIST_TEMPLATE_FILENAME,
        LOADING_SPACE_TEMPLATE_FILENAME,
        OPTIMIZATION_RESULT_TEMPLATE_FILENAME,
    }
    for path in paths:
        assert path.exists()


def test_catalog_template_imports_without_errors(tmp_path: Path) -> None:
    generate_all_templates(tmp_path, application_version="9.9.9")
    result = import_catalog(tmp_path / CATALOG_TEMPLATE_FILENAME)
    assert result.errors == ()
    assert len(result.units) >= 1


def test_loading_space_template_imports_without_errors(tmp_path: Path) -> None:
    generate_all_templates(tmp_path, application_version="9.9.9")
    result = import_loading_spaces(tmp_path / LOADING_SPACE_TEMPLATE_FILENAME)
    assert result.errors == ()
    assert len(result.spaces) >= 1


def test_packing_list_template_resolves_against_catalog_template(tmp_path: Path) -> None:
    generate_all_templates(tmp_path, application_version="9.9.9")
    catalog_result = import_catalog(tmp_path / CATALOG_TEMPLATE_FILENAME)
    catalog_by_sku = {unit.sku: unit for unit in catalog_result.units}

    packing_list_result = import_packing_list(
        tmp_path / PACKING_LIST_TEMPLATE_FILENAME, catalog_by_sku.get
    )

    assert packing_list_result.errors == ()
    assert packing_list_result.missing_skus == ()
    assert len(packing_list_result.resolved_units) >= 1


def test_optimization_result_template_has_expected_sheets(tmp_path: Path) -> None:
    generate_all_templates(tmp_path, application_version="9.9.9")
    workbook = load_workbook(tmp_path / OPTIMIZATION_RESULT_TEMPLATE_FILENAME)
    assert workbook.sheetnames == [
        "Resumen",
        "Productos cargados",
        "Productos no cargados",
        "Warnings",
        "Datos del espacio",
    ]
