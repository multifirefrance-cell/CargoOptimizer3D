"""Pruebas de `catalog_importer.py`/`catalog_exporter.py`: corruptos, duplicados, round-trip."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from cargo_optimizer.domain.dimensions import Dimensions3D
from cargo_optimizer.domain.load_unit import LoadUnit
from cargo_optimizer.infrastructure.excel.catalog_exporter import export_catalog
from cargo_optimizer.infrastructure.excel.catalog_importer import import_catalog
from cargo_optimizer.infrastructure.excel.exceptions import ExcelFileError, ExcelTemplateError
from cargo_optimizer.infrastructure.excel.product_rows import PRODUCT_COLUMNS


def _make_unit(sku: str = "SKU-1") -> LoadUnit:
    return LoadUnit(
        sku=sku,
        name="Producto de prueba",
        dimensions=Dimensions3D(50, 40, 30),
        weight_kg=10.0,
        quantity=2,
    )


def test_export_then_import_round_trips_cleanly(tmp_path: Path) -> None:
    units = (_make_unit("SKU-1"), _make_unit("SKU-2"))
    path = tmp_path / "catalogo.xlsx"

    export_catalog(units, path)
    result = import_catalog(path)

    assert result.errors == ()
    assert [u.sku for u in result.units] == ["SKU-1", "SKU-2"]
    assert result.units[0].name == units[0].name
    assert result.units[0].dimensions == units[0].dimensions


def test_import_catalog_corrupt_file_raises(tmp_path: Path) -> None:
    corrupt = tmp_path / "corrupto.xlsx"
    corrupt.write_text("no es un xlsx", encoding="utf-8")
    with pytest.raises(ExcelFileError):
        import_catalog(corrupt)


def test_import_catalog_missing_columns_raises(tmp_path: Path) -> None:
    path = tmp_path / "incompleto.xlsx"
    workbook = Workbook()
    workbook.active.append(["SKU", "Nombre"])
    workbook.active.append(["SKU-1", "Producto"])
    workbook.save(path)

    with pytest.raises(ExcelTemplateError):
        import_catalog(path)


def test_import_catalog_reports_duplicate_sku_within_file(tmp_path: Path) -> None:
    path = tmp_path / "duplicado.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(PRODUCT_COLUMNS)
    worksheet.append(
        ["SKU-1", "Producto A", 50, 40, 30, 10.0, 1, "", "", "", "", "", "", 1, "", ""]
    )
    worksheet.append(
        ["SKU-1", "Producto B", 50, 40, 30, 10.0, 1, "", "", "", "", "", "", 1, "", ""]
    )
    workbook.save(path)

    result = import_catalog(path)

    assert len(result.units) == 1
    assert result.units[0].name == "Producto A"
    assert len(result.errors) == 1
    assert result.errors[0].row_number == 3
    assert "duplicado" in result.errors[0].message.casefold()


def test_import_catalog_reports_row_with_invalid_type_without_stopping(tmp_path: Path) -> None:
    path = tmp_path / "tipo_invalido.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(PRODUCT_COLUMNS)
    worksheet.append(
        [
            "SKU-BAD",
            "Producto malo",
            "no-numerico",
            40,
            30,
            10.0,
            1,
            "",
            "",
            "",
            "",
            "",
            "",
            1,
            "",
            "",
        ]
    )
    worksheet.append(
        ["SKU-OK", "Producto bueno", 50, 40, 30, 10.0, 1, "", "", "", "", "", "", 1, "", ""]
    )
    workbook.save(path)

    result = import_catalog(path)

    assert [u.sku for u in result.units] == ["SKU-OK"]
    assert len(result.errors) == 1
    assert result.errors[0].row_number == 2


def test_import_catalog_skips_fully_blank_rows(tmp_path: Path) -> None:
    path = tmp_path / "con_fila_vacia.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(PRODUCT_COLUMNS)
    worksheet.append(["SKU-1", "Producto", 50, 40, 30, 10.0, 1, "", "", "", "", "", "", 1, "", ""])
    worksheet.append([None] * len(PRODUCT_COLUMNS))
    workbook.save(path)

    result = import_catalog(path)

    assert len(result.units) == 1
    assert result.errors == ()
