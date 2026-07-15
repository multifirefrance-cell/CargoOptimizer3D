"""Pruebas de `advanced_export.py`: selección, orden, nombre y ocultar hojas vacías (fase 8.1)."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from cargo_optimizer.domain.dimensions import Dimensions3D
from cargo_optimizer.domain.enums import LoadingSpaceCategory, OrientationCode
from cargo_optimizer.domain.load_unit import LoadUnit
from cargo_optimizer.domain.loading_space import LoadingSpace
from cargo_optimizer.domain.orientation import Orientation
from cargo_optimizer.domain.packing_result import PackingResult
from cargo_optimizer.domain.placement import Placement
from cargo_optimizer.domain.position import Position3D
from cargo_optimizer.infrastructure.excel.advanced_export import (
    SheetSelection,
    default_sheet_order,
    export_packing_result_advanced,
    is_sheet_empty,
)
from cargo_optimizer.infrastructure.excel.exceptions import ExcelError
from cargo_optimizer.infrastructure.excel.result_exporter import (
    SHEET_PACKED,
    SHEET_SPACE,
    SHEET_SUMMARY,
    SHEET_UNPACKED,
    SHEET_WARNINGS,
)


def _sample_result_with_placement() -> tuple[PackingResult, dict]:
    space = LoadingSpace(
        name="Espacio",
        category=LoadingSpaceCategory.TRUCK,
        internal_dimensions=Dimensions3D(500, 200, 200),
    )
    box = LoadUnit(
        sku="SKU-1", name="Caja", dimensions=Dimensions3D(50, 40, 30), weight_kg=10.0, quantity=1
    )
    orientation = Orientation.from_base_dimensions(box.dimensions, OrientationCode.LWH_XYZ)
    placements = (
        Placement(
            load_unit_id=box.id,
            instance_number=1,
            position=Position3D(0, 0, 0),
            orientation=orientation,
            sequence_number=1,
        ),
    )
    result = PackingResult(
        loading_space=space,
        placements=placements,
        unpacked_units=(),
        requested_count=1,
        packed_count=1,
        used_volume_cm3=placements[0].volume_cm3,
        used_weight_kg=10.0,
        execution_time_seconds=0.5,
        algorithm_name="greedy_extreme_point_v1",
    )
    return result, {box.id: box}


def test_default_sheet_order_matches_result_exporter_order() -> None:
    assert default_sheet_order() == (
        SHEET_SUMMARY,
        SHEET_PACKED,
        SHEET_UNPACKED,
        SHEET_WARNINGS,
        SHEET_SPACE,
    )


def test_is_sheet_empty_reflects_result_content() -> None:
    result, _load_units = _sample_result_with_placement()
    assert is_sheet_empty(SHEET_PACKED, result) is False
    assert is_sheet_empty(SHEET_UNPACKED, result) is True
    assert is_sheet_empty(SHEET_WARNINGS, result) is True
    assert is_sheet_empty(SHEET_SUMMARY, result) is False
    assert is_sheet_empty(SHEET_SPACE, result) is False


def test_export_advanced_default_selection_matches_full_export(tmp_path: Path) -> None:
    result, load_units_by_id = _sample_result_with_placement()
    path = tmp_path / "avanzado.xlsx"

    export_packing_result_advanced(result, load_units_by_id, path, application_version="9.9.9")

    workbook = load_workbook(path)
    assert workbook.sheetnames == list(default_sheet_order())


def test_export_advanced_can_reorder_and_rename_sheets(tmp_path: Path) -> None:
    result, load_units_by_id = _sample_result_with_placement()
    path = tmp_path / "avanzado.xlsx"
    selections = (
        SheetSelection(SHEET_SPACE, display_name="Espacio de carga"),
        SheetSelection(SHEET_SUMMARY),
    )

    export_packing_result_advanced(
        result,
        load_units_by_id,
        path,
        application_version="9.9.9",
        sheet_selections=selections,
    )

    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Espacio de carga", SHEET_SUMMARY]


def test_export_advanced_hides_empty_sheets(tmp_path: Path) -> None:
    result, load_units_by_id = _sample_result_with_placement()
    path = tmp_path / "avanzado.xlsx"

    export_packing_result_advanced(
        result, load_units_by_id, path, application_version="9.9.9", hide_empty_sheets=True
    )

    workbook = load_workbook(path)
    assert SHEET_UNPACKED not in workbook.sheetnames
    assert SHEET_WARNINGS not in workbook.sheetnames
    assert SHEET_PACKED in workbook.sheetnames


def test_export_advanced_raises_when_no_sheets_remain(tmp_path: Path) -> None:
    result, load_units_by_id = _sample_result_with_placement()
    path = tmp_path / "avanzado.xlsx"

    with pytest.raises(ExcelError):
        export_packing_result_advanced(
            result, load_units_by_id, path, application_version="9.9.9", sheet_selections=()
        )
