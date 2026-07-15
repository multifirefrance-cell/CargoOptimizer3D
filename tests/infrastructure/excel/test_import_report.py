"""Pruebas de `import_report.py`: informe individual, agregado y exportación a Excel (fase 8.1)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from cargo_optimizer.infrastructure.excel.import_plan import ImportPlan
from cargo_optimizer.infrastructure.excel.import_preview import CatalogImportPreview
from cargo_optimizer.infrastructure.excel.import_report import (
    SHEET_ERRORS,
    SHEET_SUMMARY,
    BulkImportReport,
    build_import_report,
    export_import_report,
)
from cargo_optimizer.infrastructure.excel.results import RowError


def _preview(**overrides: object) -> CatalogImportPreview:
    kwargs: dict[str, object] = {
        "row_count": 3,
        "column_count": 16,
        "new_units": (),
        "existing_units": (),
        "duplicate_errors": (),
        "invalid_errors": (),
    }
    kwargs.update(overrides)
    return CatalogImportPreview(**kwargs)  # type: ignore[arg-type]


def test_build_import_report_combines_preview_and_plan() -> None:
    preview = _preview(
        row_count=5,
        duplicate_errors=(RowError(2, "SKU duplicado."),),
        invalid_errors=(RowError(3, "Dato invalido."),),
    )
    plan = ImportPlan(to_add=(), to_update=(), ignored_skus=("IGNORED-1",))

    report = build_import_report(
        source_name="archivo.xlsx", preview=preview, plan=plan, elapsed_seconds=1.5
    )

    assert report.source_name == "archivo.xlsx"
    assert report.rows_read == 5
    assert report.imported_count == 0
    assert report.updated_count == 0
    assert report.duplicate_count == 1
    assert report.ignored_count == 1
    assert report.error_count == 2
    assert report.elapsed_seconds == 1.5


def test_bulk_import_report_totals_sum_across_reports() -> None:
    preview = _preview(row_count=2)
    plan_a = ImportPlan(to_add=("a", "b"), to_update=(), ignored_skus=())  # type: ignore[arg-type]
    plan_b = ImportPlan(to_add=(), to_update=("c",), ignored_skus=("d",))  # type: ignore[arg-type]
    report_a = build_import_report(
        source_name="a.xlsx", preview=preview, plan=plan_a, elapsed_seconds=1.0
    )
    report_b = build_import_report(
        source_name="b.xlsx", preview=preview, plan=plan_b, elapsed_seconds=2.0
    )

    bulk = BulkImportReport((report_a, report_b))

    assert bulk.total_rows_read == 4
    assert bulk.total_imported == 2
    assert bulk.total_updated == 1
    assert bulk.total_ignored == 1
    assert bulk.total_elapsed_seconds == 3.0


def test_export_import_report_writes_summary_and_error_sheets(tmp_path: Path) -> None:
    preview = _preview(row_count=2, invalid_errors=(RowError(4, "Fila invalida."),))
    plan = ImportPlan(to_add=("a",), to_update=(), ignored_skus=())  # type: ignore[arg-type]
    report = build_import_report(
        source_name="archivo.xlsx", preview=preview, plan=plan, elapsed_seconds=0.5
    )
    out_path = tmp_path / "informe.xlsx"

    export_import_report(report, out_path)

    assert out_path.exists()
    workbook = load_workbook(out_path)
    assert SHEET_SUMMARY in workbook.sheetnames
    assert SHEET_ERRORS in workbook.sheetnames
    summary_rows = list(workbook[SHEET_SUMMARY].iter_rows(values_only=True))
    assert any(row and row[0] == "archivo.xlsx" for row in summary_rows)
    error_rows = list(workbook[SHEET_ERRORS].iter_rows(values_only=True))
    assert any(row and row[0] == "archivo.xlsx" and row[1] == 4 for row in error_rows)


def test_export_import_report_bulk_adds_totals_row(tmp_path: Path) -> None:
    preview = _preview(row_count=1)
    plan = ImportPlan(to_add=("a",), to_update=(), ignored_skus=())  # type: ignore[arg-type]
    report_a = build_import_report(
        source_name="a.xlsx", preview=preview, plan=plan, elapsed_seconds=0.1
    )
    report_b = build_import_report(
        source_name="b.xlsx", preview=preview, plan=plan, elapsed_seconds=0.2
    )
    out_path = tmp_path / "informe_masivo.xlsx"

    export_import_report(BulkImportReport((report_a, report_b)), out_path)

    workbook = load_workbook(out_path)
    summary_rows = list(workbook[SHEET_SUMMARY].iter_rows(values_only=True))
    assert any(row and row[0] == "TOTAL" for row in summary_rows)
