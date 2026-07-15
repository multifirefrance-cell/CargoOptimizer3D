"""Pruebas de `loading_space_importer.py`: contenedores, camiones, vans, bodegas, personalizados."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from cargo_optimizer.domain.dimensions import Dimensions3D
from cargo_optimizer.domain.enums import DoorPosition, LoadingSpaceCategory
from cargo_optimizer.domain.loading_space import LoadingSpace
from cargo_optimizer.infrastructure.excel.exceptions import ExcelFileError, ExcelTemplateError
from cargo_optimizer.infrastructure.excel.loading_space_importer import import_loading_spaces
from cargo_optimizer.infrastructure.excel.loading_space_rows import (
    LOADING_SPACE_COLUMNS,
    loading_space_to_row,
)


def test_round_trip_of_container_truck_van_and_warehouse(tmp_path: Path) -> None:
    spaces = (
        LoadingSpace.standard_20ft_container(),
        LoadingSpace(
            name="Camion 3/4",
            category=LoadingSpaceCategory.TRUCK,
            internal_dimensions=Dimensions3D(600, 240, 240),
            door_position=DoorPosition.REAR,
            max_weight_kg=8000.0,
        ),
        LoadingSpace(
            name="Van de reparto",
            category=LoadingSpaceCategory.VAN,
            internal_dimensions=Dimensions3D(300, 170, 180),
            door_position=DoorPosition.RIGHT,
        ),
        LoadingSpace(
            name="Bodega personalizada",
            category=LoadingSpaceCategory.WAREHOUSE,
            internal_dimensions=Dimensions3D(2000, 1500, 500),
            door_position=DoorPosition.UNRESTRICTED,
            max_weight_kg=None,
        ),
    )
    path = tmp_path / "espacios.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(LOADING_SPACE_COLUMNS)
    for space in spaces:
        worksheet.append(loading_space_to_row(space))
    workbook.save(path)

    result = import_loading_spaces(path)

    assert result.errors == ()
    assert len(result.spaces) == 4
    assert result.spaces[0].category == LoadingSpaceCategory.CONTAINER
    assert result.spaces[1].category == LoadingSpaceCategory.TRUCK
    assert result.spaces[1].max_weight_kg == 8000.0
    assert result.spaces[2].category == LoadingSpaceCategory.VAN
    assert result.spaces[3].max_weight_kg is None


def test_import_loading_spaces_rejects_blank_name(tmp_path: Path) -> None:
    path = tmp_path / "espacios.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(LOADING_SPACE_COLUMNS)
    worksheet.append(["", "Contenedor", 500, 200, 200, 1000, "Trasera", ""])
    workbook.save(path)

    result = import_loading_spaces(path)

    assert result.spaces == ()
    assert len(result.errors) == 1


def test_import_loading_spaces_rejects_non_numeric_dimension(tmp_path: Path) -> None:
    path = tmp_path / "espacios.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(LOADING_SPACE_COLUMNS)
    worksheet.append(["Espacio", "Contenedor", "no-numerico", 200, 200, 1000, "Trasera", ""])
    workbook.save(path)

    result = import_loading_spaces(path)

    assert result.spaces == ()
    assert len(result.errors) == 1


def test_import_loading_spaces_missing_columns_raises(tmp_path: Path) -> None:
    path = tmp_path / "espacios.xlsx"
    workbook = Workbook()
    workbook.active.append(["Nombre"])
    workbook.save(path)

    with pytest.raises(ExcelTemplateError):
        import_loading_spaces(path)


def test_import_loading_spaces_corrupt_file_raises(tmp_path: Path) -> None:
    corrupt = tmp_path / "corrupto.xlsx"
    corrupt.write_text("no es un xlsx", encoding="utf-8")
    with pytest.raises(ExcelFileError):
        import_loading_spaces(corrupt)


def test_import_loading_spaces_defaults_blank_category_and_door(tmp_path: Path) -> None:
    path = tmp_path / "espacios.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(LOADING_SPACE_COLUMNS)
    worksheet.append(["Espacio minimo", "", 100, 100, 100, "", "", ""])
    workbook.save(path)

    result = import_loading_spaces(path)

    assert len(result.spaces) == 1
    assert result.spaces[0].category == LoadingSpaceCategory.OTHER
    assert result.spaces[0].door_position == DoorPosition.REAR
    assert result.spaces[0].max_weight_kg is None
