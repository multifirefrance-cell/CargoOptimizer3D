"""Pruebas de la plantilla oficial descargable de productos (fase OPT-18).

Distinta de `test_templates.py` (que cubre las cuatro plantillas
internas de referencia, `generate_all_templates`): aquí se prueba
`generate_product_import_template`, la plantilla vacía que se ofrece
desde "Descargar plantilla Excel…" a un usuario que todavía no conoce
la estructura del programa.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from cargo_optimizer.domain.enums import OrientationCode
from cargo_optimizer.domain.load_unit import DEFAULT_MAX_STACK_COUNT, DEFAULT_ORIENTATION_CODES
from cargo_optimizer.infrastructure.excel.catalog_importer import import_catalog
from cargo_optimizer.infrastructure.excel.product_rows import (
    EXAMPLE_ROW_SKU_MARKER,
    PRODUCT_COLUMNS,
)
from cargo_optimizer.infrastructure.excel.templates import (
    INSTRUCTIONS_SHEET_NAME,
    PRODUCT_IMPORT_SHEET_NAME,
    generate_product_import_template,
)


def test_generate_product_import_template_creates_the_two_sheets(tmp_path: Path) -> None:
    path = tmp_path / "plantilla.xlsx"
    generate_product_import_template(path)

    assert path.exists()
    workbook = load_workbook(path)
    assert workbook.sheetnames == [PRODUCT_IMPORT_SHEET_NAME, INSTRUCTIONS_SHEET_NAME]


def test_products_sheet_has_the_real_domain_headers(tmp_path: Path) -> None:
    path = tmp_path / "plantilla.xlsx"
    generate_product_import_template(path)

    workbook = load_workbook(path)
    products_sheet = workbook[PRODUCT_IMPORT_SHEET_NAME]
    header_values = tuple(
        cell.value for cell in next(products_sheet.iter_rows(min_row=1, max_row=1))
    )
    assert header_values == PRODUCT_COLUMNS


def test_instructions_sheet_documents_every_product_column(tmp_path: Path) -> None:
    path = tmp_path / "plantilla.xlsx"
    generate_product_import_template(path)

    workbook = load_workbook(path)
    instructions_sheet = workbook[INSTRUCTIONS_SHEET_NAME]
    documented_columns = {
        row[0].value for row in instructions_sheet.iter_rows(min_row=5) if row[0].value is not None
    }
    assert documented_columns == set(PRODUCT_COLUMNS)


def test_example_row_is_never_imported_as_a_real_product(tmp_path: Path) -> None:
    path = tmp_path / "plantilla.xlsx"
    generate_product_import_template(path)

    result = import_catalog(path)

    assert result.units == ()
    assert result.errors == ()
    workbook = load_workbook(path)
    products_sheet = workbook[PRODUCT_IMPORT_SHEET_NAME]
    example_sku = products_sheet.cell(row=2, column=1).value
    assert example_sku == EXAMPLE_ROW_SKU_MARKER


def test_filled_template_imports_successfully_with_real_skus(tmp_path: Path) -> None:
    path = tmp_path / "plantilla.xlsx"
    generate_product_import_template(path)
    workbook = load_workbook(path)
    products_sheet = workbook[PRODUCT_IMPORT_SHEET_NAME]
    products_sheet.append(("SKU-REAL-1", "Producto real", 30, 20, 10, 5.0, 2))
    products_sheet.append(("SKU-REAL-2", "Otro producto real", 40, 30, 20, 8.0, 1))
    workbook.save(path)

    result = import_catalog(path)

    assert result.errors == ()
    skus = {unit.sku for unit in result.units}
    assert skus == {"SKU-REAL-1", "SKU-REAL-2"}


def test_blank_optional_fields_from_the_template_get_correct_defaults(tmp_path: Path) -> None:
    path = tmp_path / "plantilla.xlsx"
    generate_product_import_template(path)
    workbook = load_workbook(path)
    products_sheet = workbook[PRODUCT_IMPORT_SHEET_NAME]
    products_sheet.append(("SKU-DEFAULTS", "Producto con campos vacios", 30, 20, 10, 5.0))
    workbook.save(path)

    result = import_catalog(path)

    assert result.errors == ()
    unit = result.units[0]
    assert unit.quantity == 1
    assert unit.max_stack_count == DEFAULT_MAX_STACK_COUNT
    assert unit.allowed_orientation_codes == DEFAULT_ORIENTATION_CODES


def test_blank_color_cells_get_distinct_pastel_colors_within_the_same_import(
    tmp_path: Path,
) -> None:
    path = tmp_path / "plantilla.xlsx"
    generate_product_import_template(path)
    workbook = load_workbook(path)
    products_sheet = workbook[PRODUCT_IMPORT_SHEET_NAME]
    products_sheet.append(("SKU-A", "Producto A", 30, 20, 10, 5.0))
    products_sheet.append(("SKU-B", "Producto B", 30, 20, 10, 5.0))
    products_sheet.append(("SKU-C", "Producto C", 30, 20, 10, 5.0))
    workbook.save(path)

    result = import_catalog(path)

    assert result.errors == ()
    colors = [unit.color_hex for unit in result.units]
    assert len(colors) == 3
    assert len(set(colors)) == 3
    for color in colors:
        assert color != "#CCCCCC"


def test_explicit_hex_color_is_preserved_through_import(tmp_path: Path) -> None:
    path = tmp_path / "plantilla.xlsx"
    generate_product_import_template(path)
    workbook = load_workbook(path)
    products_sheet = workbook[PRODUCT_IMPORT_SHEET_NAME]
    products_sheet.append(
        ("SKU-COLOR", "Producto con color explicito", 30, 20, 10, 5.0, 1, "#112233")
    )
    workbook.save(path)

    result = import_catalog(path)

    assert result.errors == ()
    assert result.units[0].color_hex == "#112233"


def test_explicit_orientations_are_preserved_through_import(tmp_path: Path) -> None:
    path = tmp_path / "plantilla.xlsx"
    generate_product_import_template(path)
    workbook = load_workbook(path)
    products_sheet = workbook[PRODUCT_IMPORT_SHEET_NAME]
    products_sheet.append(
        (
            "SKU-ORIENT",
            "Producto con orientacion explicita",
            30,
            20,
            10,
            5.0,
            1,
            "",
            "No",
            "",
            "No",
            "",
            "",
            "",
            "Todas",
        )
    )
    workbook.save(path)

    result = import_catalog(path)

    assert result.errors == ()
    assert set(result.units[0].allowed_orientation_codes) == set(OrientationCode)


def test_duplicate_real_sku_is_still_detected_alongside_the_example_row(tmp_path: Path) -> None:
    path = tmp_path / "plantilla.xlsx"
    generate_product_import_template(path)
    workbook = load_workbook(path)
    products_sheet = workbook[PRODUCT_IMPORT_SHEET_NAME]
    products_sheet.append(("SKU-DUP", "Primera aparicion", 30, 20, 10, 5.0))
    products_sheet.append(("SKU-DUP", "Segunda aparicion", 30, 20, 10, 5.0))
    workbook.save(path)

    result = import_catalog(path)

    assert len(result.units) == 1
    assert len(result.errors) == 1
    assert "duplicado" in result.errors[0].message.casefold()


def test_existing_colors_seed_is_respected_by_import_catalog(tmp_path: Path) -> None:
    """`import_catalog(..., existing_colors=...)` deja que la sugerencia de color
    pastel evite también los colores ya usados fuera de este archivo (p. ej. el
    catálogo real), no solo los vistos dentro del propio archivo."""
    path = tmp_path / "plantilla.xlsx"
    generate_product_import_template(path)
    workbook = load_workbook(path)
    products_sheet = workbook[PRODUCT_IMPORT_SHEET_NAME]
    products_sheet.append(("SKU-SEEDED", "Producto con semilla de color", 30, 20, 10, 5.0))
    workbook.save(path)

    without_seed = import_catalog(path)
    with_seed = import_catalog(path, existing_colors=(without_seed.units[0].color_hex,))

    assert without_seed.units[0].color_hex != with_seed.units[0].color_hex


def test_old_style_workbook_without_example_row_still_imports_unaffected(tmp_path: Path) -> None:
    """Compatibilidad hacia atrás: un archivo que nunca tuvo la fila de ejemplo
    reservada (p. ej. `CatalogTemplate.xlsx` u otro export anterior) sigue
    importando exactamente igual -- el marcador es puramente aditivo."""
    from cargo_optimizer.infrastructure.excel.templates import generate_catalog_template

    path = tmp_path / "catalogo_anterior.xlsx"
    generate_catalog_template(path)

    result = import_catalog(path)

    assert result.errors == ()
    assert len(result.units) == 3
    assert all(unit.sku != EXAMPLE_ROW_SKU_MARKER for unit in result.units)
