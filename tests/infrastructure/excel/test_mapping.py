"""Pruebas de `mapping.py`: detección de columnas, mapeo manual, round-trip (fase 8.1)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from cargo_optimizer.infrastructure.excel.mapping import (
    build_remapped_worksheet,
    canonical_columns_for,
    detect_column_mapping,
    import_catalog_with_mapping,
    import_loading_spaces_with_mapping,
    import_packing_list_with_mapping,
    read_source_headers,
)
from cargo_optimizer.infrastructure.excel.product_rows import PRODUCT_COLUMNS
from cargo_optimizer.infrastructure.excel.workbook_utils import get_active_worksheet, open_workbook


def _kupfer_style_file(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Código", "Descripción", "Largo", "Ancho", "Alto", "Peso bruto", "Cantidad"])
    worksheet.append(["KUP-1", "Caja Kupfer", 50, 40, 30, 12.5, 3])
    workbook.save(path)


def test_read_source_headers_reads_full_row_regardless_of_schema(tmp_path: Path) -> None:
    path = tmp_path / "kupfer.xlsx"
    _kupfer_style_file(path)

    headers = read_source_headers(path)

    assert headers == ("Código", "Descripción", "Largo", "Ancho", "Alto", "Peso bruto", "Cantidad")


def test_detect_column_mapping_recognizes_kupfer_style_aliases() -> None:
    headers = ("Código", "Descripción", "Largo", "Ancho", "Alto", "Peso bruto", "Cantidad")

    mapping, unrecognized = detect_column_mapping(headers, "catalog")

    assert mapping == {
        "Código": "SKU",
        "Descripción": "Nombre",
        "Largo": "Largo (cm)",
        "Ancho": "Ancho (cm)",
        "Alto": "Alto (cm)",
        "Peso bruto": "Peso (kg)",
        "Cantidad": "Cantidad",
    }
    assert unrecognized == ()


def test_detect_column_mapping_reports_unrecognized_headers() -> None:
    headers = ("Ref", "Producto", "Almacen", "Peso")

    mapping, unrecognized = detect_column_mapping(headers, "catalog")

    assert mapping == {"Ref": "SKU", "Producto": "Nombre", "Peso": "Peso (kg)"}
    assert unrecognized == ("Almacen",)


def test_detect_column_mapping_ignores_blank_headers() -> None:
    mapping, unrecognized = detect_column_mapping(("SKU", "", "  "), "catalog")

    assert mapping == {"SKU": "SKU"}
    assert unrecognized == ()


def test_canonical_columns_for_each_target_kind() -> None:
    assert canonical_columns_for("catalog") == PRODUCT_COLUMNS
    assert "SKU" in canonical_columns_for("packing_list")
    assert "Nombre" in canonical_columns_for("loading_space")


def test_build_remapped_worksheet_reorders_and_renames_columns(tmp_path: Path) -> None:
    path = tmp_path / "kupfer.xlsx"
    _kupfer_style_file(path)
    workbook = open_workbook(path)
    worksheet = get_active_worksheet(workbook)
    mapping, _unrecognized = detect_column_mapping(read_source_headers(path), "catalog")

    remapped = build_remapped_worksheet(worksheet, mapping, PRODUCT_COLUMNS)

    header_values = [cell.value for cell in next(remapped.iter_rows(min_row=1, max_row=1))]
    assert header_values == list(PRODUCT_COLUMNS)
    data_row = [cell.value for cell in next(remapped.iter_rows(min_row=2, max_row=2))]
    sku_index = PRODUCT_COLUMNS.index("SKU")
    name_index = PRODUCT_COLUMNS.index("Nombre")
    quantity_index = PRODUCT_COLUMNS.index("Cantidad")
    assert data_row[sku_index] == "KUP-1"
    assert data_row[name_index] == "Caja Kupfer"
    assert data_row[quantity_index] == 3


def test_build_remapped_worksheet_leaves_unmapped_canonical_columns_blank(tmp_path: Path) -> None:
    path = tmp_path / "kupfer.xlsx"
    _kupfer_style_file(path)
    workbook = open_workbook(path)
    worksheet = get_active_worksheet(workbook)
    mapping, _unrecognized = detect_column_mapping(read_source_headers(path), "catalog")

    remapped = build_remapped_worksheet(worksheet, mapping, PRODUCT_COLUMNS)

    data_row = [cell.value for cell in next(remapped.iter_rows(min_row=2, max_row=2))]
    color_index = PRODUCT_COLUMNS.index("Color")
    assert data_row[color_index] is None


def test_import_catalog_with_mapping_kupfer_style_file(tmp_path: Path) -> None:
    path = tmp_path / "kupfer.xlsx"
    _kupfer_style_file(path)
    mapping, unrecognized = detect_column_mapping(read_source_headers(path), "catalog")
    assert unrecognized == ()

    result = import_catalog_with_mapping(path, mapping)

    assert result.errors == ()
    assert len(result.units) == 1
    unit = result.units[0]
    assert unit.sku == "KUP-1"
    assert unit.name == "Caja Kupfer"
    assert unit.dimensions.length_cm == 50
    assert unit.weight_kg == 12.5
    assert unit.quantity == 3


def test_import_catalog_with_mapping_partial_mapping_reports_row_error(tmp_path: Path) -> None:
    path = tmp_path / "parcial.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Ref", "Producto", "Almacen", "Peso"])
    worksheet.append(["PART-1", "Producto parcial", "Nave 3", 5.0])
    workbook.save(path)
    mapping, unrecognized = detect_column_mapping(read_source_headers(path), "catalog")
    assert unrecognized == ("Almacen",)

    result = import_catalog_with_mapping(path, mapping)

    # Sin columnas de dimensiones mapeadas, la fila es inválida pero no
    # detiene la importación del resto del archivo.
    assert result.units == ()
    assert len(result.errors) == 1
    assert result.errors[0].row_number == 2


def test_import_packing_list_with_mapping_arbitrary_headers(tmp_path: Path) -> None:
    path = tmp_path / "packing_list_joan.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Item", "Unidades"])
    worksheet.append(["SKU-1", 5])
    workbook.save(path)
    mapping, unrecognized = detect_column_mapping(read_source_headers(path), "packing_list")
    assert unrecognized == ()

    result = import_packing_list_with_mapping(path, mapping, lambda _sku: None)

    assert result.errors == ()
    assert result.missing_skus == ("SKU-1",)


def test_import_loading_spaces_with_mapping_arbitrary_headers(tmp_path: Path) -> None:
    path = tmp_path / "espacio_exanco.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Descripción", "Largo (cm)", "Ancho (cm)", "Alto (cm)"])
    worksheet.append(["Bodega Exanco", 1200, 240, 260])
    workbook.save(path)
    mapping, unrecognized = detect_column_mapping(read_source_headers(path), "loading_space")
    assert unrecognized == ()

    result = import_loading_spaces_with_mapping(path, mapping)

    assert result.errors == ()
    assert len(result.spaces) == 1
    assert result.spaces[0].name == "Bodega Exanco"
    assert result.spaces[0].internal_dimensions.length_cm == 1200
