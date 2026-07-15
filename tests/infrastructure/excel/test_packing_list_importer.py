"""Pruebas de `packing_list_importer.py`: SKU inexistentes, duplicados, tipos incorrectos."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from cargo_optimizer.domain.dimensions import Dimensions3D
from cargo_optimizer.domain.load_unit import LoadUnit
from cargo_optimizer.infrastructure.excel.exceptions import ExcelFileError, ExcelTemplateError
from cargo_optimizer.infrastructure.excel.packing_list_importer import (
    PACKING_LIST_COLUMNS,
    import_packing_list,
)


def _catalog_resolver(units: dict[str, LoadUnit]):
    def _resolve(sku: str) -> LoadUnit | None:
        return units.get(sku)

    return _resolve


def _sample_unit(sku: str) -> LoadUnit:
    return LoadUnit(
        sku=sku, name=f"Producto {sku}", dimensions=Dimensions3D(10, 10, 10), weight_kg=1.0
    )


def test_import_packing_list_resolves_known_skus(tmp_path: Path) -> None:
    path = tmp_path / "packing_list.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(PACKING_LIST_COLUMNS)
    worksheet.append(["SKU-1", 5])
    workbook.save(path)

    catalog = {"SKU-1": _sample_unit("SKU-1")}
    result = import_packing_list(path, _catalog_resolver(catalog))

    assert result.missing_skus == ()
    assert result.errors == ()
    assert len(result.resolved_units) == 1
    assert result.resolved_units[0].sku == "SKU-1"
    assert result.resolved_units[0].quantity == 5
    # La copia resuelta es independiente del catálogo (UUID nuevo).
    assert result.resolved_units[0].id != catalog["SKU-1"].id


def test_import_packing_list_reports_missing_skus(tmp_path: Path) -> None:
    path = tmp_path / "packing_list.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(PACKING_LIST_COLUMNS)
    worksheet.append(["NO-EXISTE", 3])
    workbook.save(path)

    result = import_packing_list(path, _catalog_resolver({}))

    assert result.missing_skus == ("NO-EXISTE",)
    assert result.resolved_units == ()


def test_import_packing_list_sums_quantities_for_duplicate_sku(tmp_path: Path) -> None:
    path = tmp_path / "packing_list.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(PACKING_LIST_COLUMNS)
    worksheet.append(["SKU-1", 5])
    worksheet.append(["sku-1", 3])
    workbook.save(path)

    catalog = {"SKU-1": _sample_unit("SKU-1")}
    result = import_packing_list(path, _catalog_resolver(catalog))

    assert len(result.resolved_units) == 1
    assert result.resolved_units[0].quantity == 8


def test_import_packing_list_reports_invalid_quantity(tmp_path: Path) -> None:
    path = tmp_path / "packing_list.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(PACKING_LIST_COLUMNS)
    worksheet.append(["SKU-1", "no-numerico"])
    workbook.save(path)

    result = import_packing_list(path, _catalog_resolver({"SKU-1": _sample_unit("SKU-1")}))

    assert result.resolved_units == ()
    assert len(result.errors) == 1
    assert result.errors[0].row_number == 2


def test_import_packing_list_reports_blank_sku(tmp_path: Path) -> None:
    path = tmp_path / "packing_list.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(PACKING_LIST_COLUMNS)
    worksheet.append(["", 5])
    workbook.save(path)

    result = import_packing_list(path, _catalog_resolver({}))

    assert len(result.errors) == 1


def test_import_packing_list_missing_columns_raises(tmp_path: Path) -> None:
    path = tmp_path / "packing_list.xlsx"
    workbook = Workbook()
    workbook.active.append(["SKU"])
    workbook.save(path)

    with pytest.raises(ExcelTemplateError):
        import_packing_list(path, _catalog_resolver({}))


def test_import_packing_list_corrupt_file_raises(tmp_path: Path) -> None:
    corrupt = tmp_path / "corrupto.xlsx"
    corrupt.write_text("no es un xlsx", encoding="utf-8")
    with pytest.raises(ExcelFileError):
        import_packing_list(corrupt, _catalog_resolver({}))
