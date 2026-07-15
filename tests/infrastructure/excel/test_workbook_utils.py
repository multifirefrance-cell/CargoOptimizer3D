"""Pruebas de `infrastructure/excel/workbook_utils.py`: apertura, cabeceras, escritura atómica."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from cargo_optimizer.infrastructure.excel.exceptions import ExcelFileError, ExcelTemplateError
from cargo_optimizer.infrastructure.excel.workbook_utils import (
    get_active_worksheet,
    open_workbook,
    require_headers,
    save_workbook_atomic,
)


def test_open_workbook_missing_file_raises_excel_file_error(tmp_path: Path) -> None:
    with pytest.raises(ExcelFileError):
        open_workbook(tmp_path / "no_existe.xlsx")


def test_open_workbook_corrupt_file_raises_excel_file_error(tmp_path: Path) -> None:
    corrupt = tmp_path / "corrupto.xlsx"
    corrupt.write_text("esto no es un archivo xlsx valido", encoding="utf-8")
    with pytest.raises(ExcelFileError):
        open_workbook(corrupt)


def test_open_workbook_reads_a_real_file(tmp_path: Path) -> None:
    path = tmp_path / "real.xlsx"
    workbook = Workbook()
    workbook.active.append(["a", "b"])
    workbook.save(path)

    reopened = open_workbook(path)
    assert get_active_worksheet(reopened).cell(row=1, column=1).value == "a"


def test_get_active_worksheet_returns_a_worksheet() -> None:
    workbook = Workbook()
    worksheet = get_active_worksheet(workbook)
    assert worksheet is workbook.active


def test_require_headers_accepts_matching_headers_case_insensitive() -> None:
    workbook = Workbook()
    worksheet = get_active_worksheet(workbook)
    worksheet.append(["Sku", " Nombre "])
    require_headers(worksheet, ("SKU", "Nombre"))


def test_require_headers_rejects_missing_columns() -> None:
    workbook = Workbook()
    worksheet = get_active_worksheet(workbook)
    worksheet.append(["SKU"])
    with pytest.raises(ExcelTemplateError):
        require_headers(worksheet, ("SKU", "Nombre"))


def test_require_headers_rejects_wrong_order() -> None:
    workbook = Workbook()
    worksheet = get_active_worksheet(workbook)
    worksheet.append(["Nombre", "SKU"])
    with pytest.raises(ExcelTemplateError):
        require_headers(worksheet, ("SKU", "Nombre"))


def test_save_workbook_atomic_writes_the_file_and_leaves_no_temp_file(tmp_path: Path) -> None:
    path = tmp_path / "salida.xlsx"
    workbook = Workbook()
    workbook.active.append(["x"])

    save_workbook_atomic(workbook, path)

    assert path.exists()
    assert list(tmp_path.glob("*.tmp")) == []


def test_save_workbook_atomic_creates_parent_directory(tmp_path: Path) -> None:
    path = tmp_path / "nested" / "dir" / "salida.xlsx"
    workbook = Workbook()

    save_workbook_atomic(workbook, path)

    assert path.exists()
