"""Pruebas de integración: `MainWindow` y `infrastructure/excel` (fase 8.0).

Usa archivos `.xlsx` reales (generados con las propias plantillas
oficiales, `tmp_path`) y un `CatalogService` real respaldado por SQLite
temporal — mismo criterio que la suite de catálogo de la fase 7.1:
preferir el camino real sobre mocks siempre que sea seguro y rápido
bajo `offscreen`. Solo se sustituyen los diálogos nativos de Qt
(`QFileDialog`, `QMessageBox`, `QInputDialog`), nunca el motor de
importación/exportación real.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook
from PySide6.QtWidgets import QApplication, QFileDialog, QInputDialog, QMessageBox

from cargo_optimizer.domain.dimensions import Dimensions3D
from cargo_optimizer.domain.enums import LoadingSpaceCategory
from cargo_optimizer.domain.load_unit import LoadUnit
from cargo_optimizer.domain.loading_space import LoadingSpace
from cargo_optimizer.domain.packing_result import PackingResult
from cargo_optimizer.infrastructure.database.catalog_service import CatalogService
from cargo_optimizer.infrastructure.excel.catalog_exporter import export_catalog
from cargo_optimizer.infrastructure.excel.loading_space_rows import (
    LOADING_SPACE_COLUMNS,
    loading_space_to_row,
)
from cargo_optimizer.infrastructure.excel.packing_list_importer import PACKING_LIST_COLUMNS
from cargo_optimizer.infrastructure.excel.templates import generate_all_templates
from cargo_optimizer.presentation.desktop.main_window import MainWindow
from cargo_optimizer.presentation.desktop.settings import AppSettings


def _unit(**overrides: object) -> LoadUnit:
    kwargs: dict[str, object] = {
        "sku": "BOX-1",
        "name": "Caja",
        "dimensions": Dimensions3D(40.0, 30.0, 20.0),
        "weight_kg": 10.0,
    }
    kwargs.update(overrides)
    return LoadUnit(**kwargs)  # type: ignore[arg-type]


def _fake_open_dialog(path: Path | None) -> Any:
    text = str(path) if path is not None else ""
    return staticmethod(lambda *a, **k: (text, ""))


def _fake_save_dialog(path: Path | None) -> Any:
    text = str(path) if path is not None else ""
    return staticmethod(lambda *a, **k: (text, ""))


def test_import_excel_auto_dispatches_catalog_shaped_file_into_project(
    qapp: QApplication, app_settings: AppSettings, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    paths = generate_all_templates(tmp_path, application_version="9.9.9")
    catalog_path = next(p for p in paths if p.name == "CatalogTemplate.xlsx")
    window = MainWindow(app_settings, catalog_service=None)
    monkeypatch.setattr(QFileDialog, "getOpenFileName", _fake_open_dialog(catalog_path))

    window._on_import_excel_auto()

    assert window.product_table_panel.model.rowCount() >= 1
    assert window._is_dirty is True
    window.close()


def test_import_excel_auto_dispatches_loading_space_shaped_file(
    qapp: QApplication, app_settings: AppSettings, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    paths = generate_all_templates(tmp_path, application_version="9.9.9")
    space_path = next(p for p in paths if p.name == "LoadingSpaceTemplate.xlsx")
    window = MainWindow(app_settings, catalog_service=None)
    monkeypatch.setattr(QFileDialog, "getOpenFileName", _fake_open_dialog(space_path))
    # La plantilla trae varios espacios de ejemplo: se necesita elegir uno.
    monkeypatch.setattr(QInputDialog, "getItem", staticmethod(lambda *a, **k: (a[3][0], True)))

    window._on_import_excel_auto()

    built = window.loading_space_form_panel.build_loading_space()
    assert built is not None
    window.close()


def test_import_excel_auto_warns_on_unrecognized_file(
    qapp: QApplication, app_settings: AppSettings, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "desconocido.xlsx"
    workbook = Workbook()
    workbook.active.append(["Columna A", "Columna B"])
    workbook.save(path)
    window = MainWindow(app_settings, catalog_service=None)
    monkeypatch.setattr(QFileDialog, "getOpenFileName", _fake_open_dialog(path))
    warnings: list[tuple[Any, ...]] = []
    monkeypatch.setattr(QMessageBox, "warning", staticmethod(lambda *a, **k: warnings.append(a)))

    window._on_import_excel_auto()

    assert len(warnings) == 1
    window.close()


def test_import_excel_auto_shows_error_on_corrupt_file(
    qapp: QApplication, app_settings: AppSettings, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    corrupt = tmp_path / "corrupto.xlsx"
    corrupt.write_text("no es un xlsx", encoding="utf-8")
    window = MainWindow(app_settings, catalog_service=None)
    monkeypatch.setattr(QFileDialog, "getOpenFileName", _fake_open_dialog(corrupt))
    errors: list[tuple[Any, ...]] = []
    monkeypatch.setattr(QMessageBox, "critical", staticmethod(lambda *a, **k: errors.append(a)))

    window._on_import_excel_auto()

    assert len(errors) == 1
    window.close()


def test_on_import_catalog_excel_adds_to_sqlite_catalog(
    qapp: QApplication, app_settings: AppSettings, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    service = CatalogService.create_default(base_dir=tmp_path / "db")
    window = MainWindow(app_settings, catalog_service=service)
    excel_path = tmp_path / "catalogo.xlsx"
    export_catalog([_unit(sku="CAT-EXCEL-1"), _unit(sku="CAT-EXCEL-2")], excel_path)
    monkeypatch.setattr(QFileDialog, "getOpenFileName", _fake_open_dialog(excel_path))

    window._on_import_catalog_excel()

    assert service.products.get_by_sku("CAT-EXCEL-1") is not None
    assert service.products.get_by_sku("CAT-EXCEL-2") is not None
    # No debe haber tocado el proyecto: catálogo y proyecto son destinos distintos.
    assert window.product_table_panel.model.rowCount() == 0
    window.close()


def test_on_import_catalog_excel_skips_existing_sku_with_warning(
    qapp: QApplication, app_settings: AppSettings, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    service = CatalogService.create_default(base_dir=tmp_path / "db")
    service.products.add(_unit(sku="ALREADY-THERE"))
    window = MainWindow(app_settings, catalog_service=service)
    excel_path = tmp_path / "catalogo.xlsx"
    export_catalog([_unit(sku="ALREADY-THERE", weight_kg=99.0)], excel_path)
    monkeypatch.setattr(QFileDialog, "getOpenFileName", _fake_open_dialog(excel_path))
    warnings: list[tuple[Any, ...]] = []
    monkeypatch.setattr(QMessageBox, "warning", staticmethod(lambda *a, **k: warnings.append(a)))

    window._on_import_catalog_excel()

    assert service.products.get_by_sku("ALREADY-THERE").weight_kg == 10.0  # type: ignore[union-attr]
    assert len(warnings) == 1
    window.close()


def test_on_import_catalog_excel_disabled_in_limited_mode(
    qapp: QApplication, app_settings: AppSettings
) -> None:
    window = MainWindow(app_settings, catalog_service=None)
    assert window.action_import_catalog_excel.isEnabled() is False
    assert window.action_export_catalog_excel.isEnabled() is False
    assert window.action_import_packing_list_excel.isEnabled() is False
    window.close()


def test_on_export_catalog_excel_writes_active_products(
    qapp: QApplication, app_settings: AppSettings, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    service = CatalogService.create_default(base_dir=tmp_path / "db")
    service.products.add(_unit(sku="EXPORT-1"))
    window = MainWindow(app_settings, catalog_service=service)
    out_path = tmp_path / "salida_catalogo.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", _fake_save_dialog(out_path))

    window._on_export_catalog_excel()

    assert out_path.exists()
    window.close()


def test_on_export_catalog_excel_warns_when_catalog_empty(
    qapp: QApplication, app_settings: AppSettings, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    service = CatalogService.create_default(base_dir=tmp_path / "db")
    window = MainWindow(app_settings, catalog_service=service)
    warnings: list[tuple[Any, ...]] = []
    monkeypatch.setattr(QMessageBox, "warning", staticmethod(lambda *a, **k: warnings.append(a)))

    window._on_export_catalog_excel()

    assert len(warnings) == 1
    window.close()


def test_on_import_packing_list_excel_adds_resolved_units(
    qapp: QApplication, app_settings: AppSettings, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    service = CatalogService.create_default(base_dir=tmp_path / "db")
    service.products.add(_unit(sku="PL-1"))
    window = MainWindow(app_settings, catalog_service=service)
    path = tmp_path / "packing_list.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(PACKING_LIST_COLUMNS)
    worksheet.append(["PL-1", 7])
    workbook.save(path)
    monkeypatch.setattr(QFileDialog, "getOpenFileName", _fake_open_dialog(path))

    window._on_import_packing_list_excel()

    units = window.product_table_panel.model.load_units()
    assert len(units) == 1
    assert units[0].sku == "PL-1"
    assert units[0].quantity == 7
    window.close()


def test_on_import_packing_list_excel_missing_sku_cancel_imports_nothing(
    qapp: QApplication, app_settings: AppSettings, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    service = CatalogService.create_default(base_dir=tmp_path / "db")
    window = MainWindow(app_settings, catalog_service=service)
    path = tmp_path / "packing_list.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(PACKING_LIST_COLUMNS)
    worksheet.append(["NO-EXISTE", 3])
    workbook.save(path)
    monkeypatch.setattr(QFileDialog, "getOpenFileName", _fake_open_dialog(path))
    monkeypatch.setattr(
        QMessageBox, "question", staticmethod(lambda *a, **k: QMessageBox.StandardButton.Cancel)
    )

    window._on_import_packing_list_excel()

    assert window.product_table_panel.model.rowCount() == 0
    window.close()


def test_on_import_packing_list_excel_missing_sku_continue_imports_the_rest(
    qapp: QApplication, app_settings: AppSettings, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    service = CatalogService.create_default(base_dir=tmp_path / "db")
    service.products.add(_unit(sku="FOUND-1"))
    window = MainWindow(app_settings, catalog_service=service)
    path = tmp_path / "packing_list.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(PACKING_LIST_COLUMNS)
    worksheet.append(["FOUND-1", 2])
    worksheet.append(["NO-EXISTE", 3])
    workbook.save(path)
    monkeypatch.setattr(QFileDialog, "getOpenFileName", _fake_open_dialog(path))
    monkeypatch.setattr(
        QMessageBox, "question", staticmethod(lambda *a, **k: QMessageBox.StandardButton.Yes)
    )

    window._on_import_packing_list_excel()

    units = window.product_table_panel.model.load_units()
    assert len(units) == 1
    assert units[0].sku == "FOUND-1"
    window.close()


def test_on_export_result_excel_warns_when_no_result(
    qapp: QApplication, app_settings: AppSettings, monkeypatch: pytest.MonkeyPatch
) -> None:
    window = MainWindow(app_settings, catalog_service=None)
    warnings: list[tuple[Any, ...]] = []
    monkeypatch.setattr(QMessageBox, "warning", staticmethod(lambda *a, **k: warnings.append(a)))

    window._on_export_result_excel()

    assert len(warnings) == 1
    window.close()


def test_on_export_result_excel_writes_file_when_result_present(
    qapp: QApplication, app_settings: AppSettings, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    window = MainWindow(app_settings, catalog_service=None)
    window._last_result = PackingResult(
        loading_space=LoadingSpace(
            name="Espacio",
            category=LoadingSpaceCategory.WAREHOUSE,
            internal_dimensions=Dimensions3D(100, 100, 100),
        ),
        placements=(),
        unpacked_units=(),
        requested_count=0,
        packed_count=0,
        used_volume_cm3=0.0,
        used_weight_kg=0.0,
        execution_time_seconds=0.01,
        algorithm_name="greedy_extreme_point_v1",
    )
    out_path = tmp_path / "resultado.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", _fake_save_dialog(out_path))

    window._on_export_result_excel()

    assert out_path.exists()
    window.close()


def test_on_import_loading_space_excel_applies_single_space_directly(
    qapp: QApplication, app_settings: AppSettings, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    window = MainWindow(app_settings, catalog_service=None)
    space = LoadingSpace(
        name="Espacio unico de prueba",
        category=LoadingSpaceCategory.VAN,
        internal_dimensions=Dimensions3D(300, 170, 180),
    )
    path = tmp_path / "espacio.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(LOADING_SPACE_COLUMNS)
    worksheet.append(loading_space_to_row(space))
    workbook.save(path)
    monkeypatch.setattr(QFileDialog, "getOpenFileName", _fake_open_dialog(path))

    window._on_import_loading_space_excel()

    built = window.loading_space_form_panel.build_loading_space()
    assert built is not None
    assert built.name == "Espacio unico de prueba"
    window.close()


def test_on_import_loading_space_excel_prompts_when_multiple_spaces(
    qapp: QApplication, app_settings: AppSettings, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    window = MainWindow(app_settings, catalog_service=None)
    spaces = (
        LoadingSpace.standard_20ft_container(),
        LoadingSpace.standard_40ft_container(),
    )
    path = tmp_path / "espacios.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(LOADING_SPACE_COLUMNS)
    for space in spaces:
        worksheet.append(loading_space_to_row(space))
    workbook.save(path)
    monkeypatch.setattr(QFileDialog, "getOpenFileName", _fake_open_dialog(path))
    monkeypatch.setattr(
        QInputDialog,
        "getItem",
        staticmethod(lambda *a, **k: (spaces[1].name, True)),
    )

    window._on_import_loading_space_excel()

    built = window.loading_space_form_panel.build_loading_space()
    assert built is not None
    assert built.name == spaces[1].name
    window.close()
