"""Pruebas de `BulkImportDialog`: varios archivos, resumen final, errores por archivo (fase 8.1)."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication, QDialogButtonBox, QFileDialog, QMessageBox

from cargo_optimizer.infrastructure.excel.exceptions import ExcelError
from cargo_optimizer.infrastructure.excel.import_report import ImportReport
from cargo_optimizer.presentation.desktop.dialogs.bulk_import_dialog import BulkImportDialog


def _report(name: str, *, imported: int = 1, errors: int = 0) -> ImportReport:
    return ImportReport(
        source_name=name,
        rows_read=imported + errors,
        imported_count=imported,
        updated_count=0,
        duplicate_count=0,
        ignored_count=0,
        errors=(),
        elapsed_seconds=0.1,
    )


def test_import_all_calls_import_one_for_every_selected_file(
    qapp: QApplication, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    paths = [tmp_path / "a.xlsx", tmp_path / "b.xlsx"]
    for path in paths:
        path.write_text("", encoding="utf-8")
    called: list[Path] = []

    def _import_one(path: Path) -> ImportReport:
        called.append(path)
        return _report(path.name)

    dialog = BulkImportDialog(None, import_one=_import_one)
    monkeypatch.setattr(
        QFileDialog, "getOpenFileNames", staticmethod(lambda *a, **k: ([str(p) for p in paths], ""))
    )
    dialog._on_choose_files()

    dialog._on_import_all()

    assert called == paths
    assert [r.source_name for r in dialog.reports()] == ["a.xlsx", "b.xlsx"]


def test_import_all_continues_after_one_file_fails(
    qapp: QApplication, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    paths = [tmp_path / "bad.xlsx", tmp_path / "good.xlsx"]

    def _import_one(path: Path) -> ImportReport:
        if path.name == "bad.xlsx":
            raise ExcelError("archivo corrupto")
        return _report(path.name)

    dialog = BulkImportDialog(None, import_one=_import_one)
    monkeypatch.setattr(
        QFileDialog, "getOpenFileNames", staticmethod(lambda *a, **k: ([str(p) for p in paths], ""))
    )
    monkeypatch.setattr(QMessageBox, "warning", staticmethod(lambda *a, **k: None))
    dialog._on_choose_files()

    dialog._on_import_all()

    assert len(dialog.reports()) == 2
    assert dialog.reports()[0].source_name == "bad.xlsx"
    assert dialog.reports()[0].imported_count == 0
    # El fallo debe quedar reflejado en el informe (no solo en el
    # QMessageBox del momento): un archivo que no se pudo leer nunca
    # puede reportar error_count == 0, igual que uno vacío pero válido.
    assert dialog.reports()[0].error_count == 1
    assert dialog.reports()[1].source_name == "good.xlsx"
    assert dialog.reports()[1].imported_count == 1


def test_close_button_only_rejects_once(
    qapp: QApplication, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    dialog = BulkImportDialog(None, import_one=lambda path: _report(path.name))
    accept_calls = []
    reject_calls = []
    monkeypatch.setattr(dialog, "accept", lambda: accept_calls.append(1))
    monkeypatch.setattr(dialog, "reject", lambda: reject_calls.append(1))

    close_box = dialog.findChild(QDialogButtonBox)
    assert close_box is not None
    close_box.button(QDialogButtonBox.StandardButton.Close).click()

    assert reject_calls == [1]
    assert accept_calls == []


def test_save_report_writes_a_summary_workbook(
    qapp: QApplication, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    dialog = BulkImportDialog(None, import_one=lambda path: _report(path.name))
    monkeypatch.setattr(
        QFileDialog,
        "getOpenFileNames",
        staticmethod(lambda *a, **k: ([str(tmp_path / "a.xlsx")], "")),
    )
    dialog._on_choose_files()
    dialog._on_import_all()
    out_path = tmp_path / "informe.xlsx"
    monkeypatch.setattr(
        QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: (str(out_path), ""))
    )
    monkeypatch.setattr(QMessageBox, "information", staticmethod(lambda *a, **k: None))

    dialog._on_save_report()

    assert out_path.exists()
    workbook = load_workbook(out_path)
    assert "Resumen" in workbook.sheetnames
