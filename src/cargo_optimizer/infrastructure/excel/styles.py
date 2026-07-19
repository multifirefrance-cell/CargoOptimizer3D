"""Formato profesional compartido para todas las hojas de `infrastructure/excel/`.

Sin macros, sin VBA — solo estilos nativos de `openpyxl` (fuentes,
rellenos, bordes) y tablas nativas de Excel (`openpyxl.worksheet.table.Table`).
Centralizado aquí para que ningún importador/exportador reimplemente su
propio formato ad-hoc.
"""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FILL = PatternFill(start_color="FFDCE6F1", end_color="FFDCE6F1", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FF1F3864")
TITLE_FONT = Font(bold=True, size=13, color="FF1F3864")
LABEL_FONT = Font(bold=True)

_THIN_SIDE = Side(style="thin", color="FFB0B0B0")
THIN_BORDER = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)

_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
_MIN_COLUMN_WIDTH = 10.0
_MAX_COLUMN_WIDTH = 60.0
_COLUMN_WIDTH_PADDING = 2.0


def write_header_row(worksheet: Worksheet, headers: tuple[str, ...], *, row: int = 1) -> None:
    """Escribe una fila de cabeceras con negrita, relleno suave, bordes y ajuste de texto."""
    for column_index, title in enumerate(headers, start=1):
        cell = worksheet.cell(row=row, column=column_index, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = _HEADER_ALIGNMENT


def apply_borders_to_data_rows(
    worksheet: Worksheet, *, first_row: int, last_row: int, column_count: int
) -> None:
    """Aplica el mismo borde fino a todas las celdas de datos (sin tocar el formato de valor)."""
    for row in range(first_row, last_row + 1):
        for column in range(1, column_count + 1):
            worksheet.cell(row=row, column=column).border = THIN_BORDER


def autofit_columns(worksheet: Worksheet, *, column_count: int) -> None:
    """Aproxima un autoajuste de ancho de columna a partir del contenido real.

    `openpyxl` no soporta autoajuste nativo (depende del renderizado de
    Excel, que no existe al escribir el archivo): se estima el ancho a
    partir de la longitud de texto más larga de cada columna, con un
    mínimo y un máximo razonables para evitar columnas ilegibles o
    absurdamente anchas.
    """
    for column_index in range(1, column_count + 1):
        letter = get_column_letter(column_index)
        longest = 0
        for cell in worksheet[letter]:
            if cell.value is None:
                continue
            longest = max(longest, len(str(cell.value)))
        width = min(max(longest + _COLUMN_WIDTH_PADDING, _MIN_COLUMN_WIDTH), _MAX_COLUMN_WIDTH)
        worksheet.column_dimensions[letter].width = width


def apply_table(
    worksheet: Worksheet,
    *,
    table_name: str,
    first_row: int,
    last_row: int,
    column_count: int,
) -> None:
    """Convierte un rango en una tabla nativa de Excel (con filtro y estilo suave).

    ``table_name`` debe ser único dentro del libro y sin espacios (lo
    exige el formato `.xlsx`). Si no hay ninguna fila de datos
    (``last_row < first_row``), no se crea tabla: una tabla de Excel
    exige al menos una fila además de la cabecera.
    """
    if last_row < first_row:
        return
    last_column_letter = get_column_letter(column_count)
    ref = f"A{first_row - 1}:{last_column_letter}{last_row}"
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showRowStripes=True,
        showFirstColumn=False,
        showLastColumn=False,
        showColumnStripes=False,
    )
    worksheet.add_table(table)
