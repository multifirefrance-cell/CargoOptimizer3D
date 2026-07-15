"""Informe final de una importación (fase 8.1): contadores + tiempo, mostrado y exportable.

Un `ImportReport` resume una única importación (un archivo); un
`BulkImportReport` agrupa varios (importación masiva, sección 8 de la
fase 8.1) y añade los totales. Ambos se pueden exportar a `.xlsx` con
el mismo formato profesional que el resto del módulo.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from cargo_optimizer.infrastructure.excel.import_plan import ImportPlan
from cargo_optimizer.infrastructure.excel.import_preview import CatalogImportPreview
from cargo_optimizer.infrastructure.excel.results import RowError
from cargo_optimizer.infrastructure.excel.styles import (
    LABEL_FONT,
    TITLE_FONT,
    apply_borders_to_data_rows,
    apply_table,
    autofit_columns,
    write_header_row,
)
from cargo_optimizer.infrastructure.excel.workbook_utils import (
    HEADER_ROW,
    get_active_worksheet,
    save_workbook_atomic,
)

SHEET_SUMMARY = "Resumen"
SHEET_ERRORS = "Errores"

_ERROR_COLUMNS = ("Archivo", "Fila", "Mensaje")


@dataclass(frozen=True, slots=True)
class ImportReport:
    """Resultado final de importar un único archivo."""

    source_name: str
    rows_read: int
    imported_count: int
    updated_count: int
    duplicate_count: int
    ignored_count: int
    errors: tuple[RowError, ...]
    elapsed_seconds: float

    @property
    def error_count(self) -> int:
        return len(self.errors)


def build_import_report(
    *,
    source_name: str,
    preview: CatalogImportPreview,
    plan: ImportPlan,
    elapsed_seconds: float,
) -> ImportReport:
    """Construye el informe final combinando la vista previa y el plan ya aplicado."""
    return ImportReport(
        source_name=source_name,
        rows_read=preview.row_count,
        imported_count=len(plan.to_add),
        updated_count=len(plan.to_update),
        duplicate_count=preview.duplicate_count,
        ignored_count=len(plan.ignored_skus),
        errors=preview.invalid_errors + preview.duplicate_errors,
        elapsed_seconds=elapsed_seconds,
    )


@dataclass(frozen=True, slots=True)
class BulkImportReport:
    """Informe agregado de una importación masiva (varios archivos, sección 8)."""

    reports: tuple[ImportReport, ...]

    @property
    def total_rows_read(self) -> int:
        return sum(report.rows_read for report in self.reports)

    @property
    def total_imported(self) -> int:
        return sum(report.imported_count for report in self.reports)

    @property
    def total_updated(self) -> int:
        return sum(report.updated_count for report in self.reports)

    @property
    def total_duplicate(self) -> int:
        return sum(report.duplicate_count for report in self.reports)

    @property
    def total_ignored(self) -> int:
        return sum(report.ignored_count for report in self.reports)

    @property
    def total_errors(self) -> int:
        return sum(report.error_count for report in self.reports)

    @property
    def total_elapsed_seconds(self) -> float:
        return sum(report.elapsed_seconds for report in self.reports)


def _write_summary_rows(worksheet: Worksheet, reports: tuple[ImportReport, ...]) -> None:
    worksheet.cell(row=1, column=1, value="Informe de importación").font = TITLE_FONT
    headers = (
        "Archivo",
        "Filas leídas",
        "Importadas",
        "Actualizadas",
        "Duplicadas",
        "Ignoradas",
        "Errores",
        "Tiempo (s)",
    )
    write_header_row(worksheet, headers, row=3)
    last_row = 3
    for report in reports:
        last_row += 1
        worksheet.append(
            (
                report.source_name,
                report.rows_read,
                report.imported_count,
                report.updated_count,
                report.duplicate_count,
                report.ignored_count,
                report.error_count,
                round(report.elapsed_seconds, 3),
            )
        )
    if len(reports) > 1:
        last_row += 1
        totals = BulkImportReport(reports)
        totals_row = (
            "TOTAL",
            totals.total_rows_read,
            totals.total_imported,
            totals.total_updated,
            totals.total_duplicate,
            totals.total_ignored,
            totals.total_errors,
            round(totals.total_elapsed_seconds, 3),
        )
        worksheet.append(totals_row)
        for column in range(1, len(headers) + 1):
            worksheet.cell(row=last_row, column=column).font = LABEL_FONT

    column_count = len(headers)
    apply_borders_to_data_rows(worksheet, first_row=3, last_row=last_row, column_count=column_count)
    if last_row > 4:
        apply_table(
            worksheet,
            table_name="TablaResumenImportacion",
            first_row=4,
            last_row=last_row,
            column_count=column_count,
        )
    autofit_columns(worksheet, column_count=column_count)


def _write_error_rows(worksheet: Worksheet, reports: tuple[ImportReport, ...]) -> None:
    write_header_row(worksheet, _ERROR_COLUMNS, row=HEADER_ROW)
    last_row = HEADER_ROW
    for report in reports:
        for error in report.errors:
            last_row += 1
            worksheet.append((report.source_name, error.row_number, error.message))
    column_count = len(_ERROR_COLUMNS)
    apply_borders_to_data_rows(
        worksheet, first_row=HEADER_ROW, last_row=last_row, column_count=column_count
    )
    if last_row > HEADER_ROW:
        apply_table(
            worksheet,
            table_name="TablaErroresImportacion",
            first_row=HEADER_ROW + 1,
            last_row=last_row,
            column_count=column_count,
        )
    autofit_columns(worksheet, column_count=column_count)


def export_import_report(report: ImportReport | BulkImportReport, path: Path) -> None:
    """Exporta un informe (de un archivo o de una importación masiva completa) a `.xlsx`."""
    reports = report.reports if isinstance(report, BulkImportReport) else (report,)

    workbook = Workbook()
    summary_sheet = get_active_worksheet(workbook)
    summary_sheet.title = SHEET_SUMMARY
    _write_summary_rows(summary_sheet, reports)

    errors_sheet = workbook.create_sheet(SHEET_ERRORS)
    _write_error_rows(errors_sheet, reports)

    save_workbook_atomic(workbook, path)
