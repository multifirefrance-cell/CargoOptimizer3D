"""Mapeador de columnas (fase 8.1): reduce al mínimo el trabajo manual al importar.

Un archivo real de logística casi nunca usa exactamente las cabeceras
oficiales de `CatalogTemplate.xlsx`/`PackingListTemplate.xlsx`/
`LoadingSpaceTemplate.xlsx`. Este módulo detecta automáticamente qué
columna del archivo corresponde a qué campo del sistema (por alias
conocidos), permite completar manualmente las que no reconoce, y
reconstruye una hoja "canónica" en memoria — reordenada y renombrada —
para poder reutilizar sin duplicar el bucle de conversión fila a fila
ya existente en `catalog_importer.py`/`packing_list_importer.py`/
`loading_space_importer.py` (fase 8.0).

No depende de `infrastructure/database`: los perfiles de mapeo
guardados (`ImportMappingProfileRepository`) se leen en
`presentation/desktop/main_window.py` y se pasan aquí ya como un
`Mapping[str, str]` plano — mismo criterio que `packing_list_importer.py`
evitando depender de la base de datos directamente.
"""

from __future__ import annotations

from collections.abc import Callable, Mapping, Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from cargo_optimizer.domain.load_unit import LoadUnit
from cargo_optimizer.infrastructure.excel.catalog_importer import import_catalog_from_worksheet
from cargo_optimizer.infrastructure.excel.detection import TemplateKind
from cargo_optimizer.infrastructure.excel.loading_space_importer import (
    import_loading_spaces_from_worksheet,
)
from cargo_optimizer.infrastructure.excel.loading_space_rows import LOADING_SPACE_COLUMNS
from cargo_optimizer.infrastructure.excel.packing_list_importer import (
    PACKING_LIST_COLUMNS,
    import_packing_list_from_worksheet,
)
from cargo_optimizer.infrastructure.excel.product_rows import PRODUCT_COLUMNS
from cargo_optimizer.infrastructure.excel.results import (
    CatalogImportResult,
    LoadingSpaceImportResult,
    PackingListImportResult,
)
from cargo_optimizer.infrastructure.excel.workbook_utils import (
    HEADER_ROW,
    get_active_worksheet,
    open_workbook,
    read_header_row,
)

CATALOG_COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "SKU": ("SKU", "Codigo", "Código", "Cod", "Referencia", "Ref", "Item", "Código de producto"),
    "Nombre": ("Nombre", "Descripcion", "Descripción", "Producto", "Denominación", "Denominacion"),
    "Largo (cm)": ("Largo (cm)", "Largo", "Largo cm", "Length", "Length (cm)"),
    "Ancho (cm)": ("Ancho (cm)", "Ancho", "Ancho cm", "Width", "Width (cm)"),
    "Alto (cm)": ("Alto (cm)", "Alto", "Alto cm", "Height", "Height (cm)"),
    "Peso (kg)": ("Peso (kg)", "Peso", "Peso bruto", "Peso Kg", "Weight", "Weight (kg)"),
    "Cantidad": ("Cantidad", "Unidades", "Qty", "Quantity"),
    "Color": ("Color",),
    "Fragil": ("Fragil", "Frágil", "Fragile"),
    "Tipo de empaque": ("Tipo de empaque", "Empaque", "Package Type"),
    "Extintor": ("Extintor", "Es extintor"),
    "Agente": ("Agente", "Agente extintor"),
    "Peso nominal (kg)": ("Peso nominal (kg)", "Peso nominal"),
    "Apilamiento": ("Apilamiento", "Max apilamiento", "Stacking"),
    "Orientaciones": ("Orientaciones", "Orientacion", "Orientación"),
    "Notas": ("Notas", "Observaciones", "Comentarios"),
}

PACKING_LIST_COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "SKU": CATALOG_COLUMN_ALIASES["SKU"],
    "Cantidad": CATALOG_COLUMN_ALIASES["Cantidad"],
}

LOADING_SPACE_COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "Nombre": ("Nombre", "Descripcion", "Descripción"),
    "Categoria": ("Categoria", "Categoría", "Tipo"),
    "Largo (cm)": CATALOG_COLUMN_ALIASES["Largo (cm)"],
    "Ancho (cm)": CATALOG_COLUMN_ALIASES["Ancho (cm)"],
    "Alto (cm)": CATALOG_COLUMN_ALIASES["Alto (cm)"],
    "Peso maximo (kg)": ("Peso maximo (kg)", "Peso máximo (kg)", "Peso maximo", "Max weight"),
    "Posicion de puerta": ("Posicion de puerta", "Posición de puerta", "Puerta"),
    "Notas": ("Notas", "Observaciones"),
}

_SCHEMAS: dict[TemplateKind, tuple[tuple[str, ...], dict[str, tuple[str, ...]]]] = {
    "catalog": (PRODUCT_COLUMNS, CATALOG_COLUMN_ALIASES),
    "packing_list": (PACKING_LIST_COLUMNS, PACKING_LIST_COLUMN_ALIASES),
    "loading_space": (LOADING_SPACE_COLUMNS, LOADING_SPACE_COLUMN_ALIASES),
}


def canonical_columns_for(target_kind: TemplateKind) -> tuple[str, ...]:
    """Las columnas canónicas del esquema de destino (para mostrarlas en el asistente)."""
    return _SCHEMAS[target_kind][0]


def read_source_headers(path: Path) -> tuple[str, ...]:
    """Lee la fila de cabecera completa de un archivo, sin asumir ningún esquema."""
    workbook = open_workbook(path)
    worksheet = get_active_worksheet(workbook)
    return read_header_row(worksheet)


def detect_column_mapping(
    headers: Sequence[str], target_kind: TemplateKind
) -> tuple[dict[str, str], tuple[str, ...]]:
    """Detecta automáticamente qué cabeceras del archivo son columnas conocidas.

    Devuelve ``(mapeo, no_reconocidas)``: ``mapeo`` es
    ``{cabecera_del_archivo: columna_canonica}`` solo para las
    cabeceras que se pudieron identificar por alias; ``no_reconocidas``
    son las que no coinciden con ningún alias conocido y requieren
    asignación manual en el asistente.
    """
    _columns, aliases = _SCHEMAS[target_kind]
    mapping: dict[str, str] = {}
    unrecognized: list[str] = []
    for header in headers:
        if not header.strip():
            continue
        normalized = header.strip().casefold()
        matched_canonical: str | None = None
        for canonical, alias_list in aliases.items():
            if any(normalized == alias.casefold() for alias in alias_list):
                matched_canonical = canonical
                break
        if matched_canonical is not None:
            mapping[header] = matched_canonical
        else:
            unrecognized.append(header)
    return mapping, tuple(unrecognized)


def build_remapped_worksheet(
    source_worksheet: Worksheet,
    column_mapping: Mapping[str, str],
    canonical_columns: tuple[str, ...],
) -> Worksheet:
    """Devuelve una hoja nueva, en memoria, reordenada y renombrada según ``column_mapping``.

    ``column_mapping`` es ``{cabecera_del_archivo: columna_canonica}``.
    Las columnas canónicas sin ninguna cabecera de origen asignada
    quedan en blanco en la hoja resultante — los importadores ya
    existentes las tratan como valores vacíos (aplican su valor por
    defecto o las rechazan si son obligatorias), sin necesitar ninguna
    lógica nueva.
    """
    source_headers = read_header_row(source_worksheet)
    source_column_by_canonical: dict[str, int] = {}
    for index, header in enumerate(source_headers, start=1):
        canonical = column_mapping.get(header.strip())
        if canonical is not None and canonical not in source_column_by_canonical:
            source_column_by_canonical[canonical] = index

    remapped_workbook = Workbook()
    remapped_worksheet = get_active_worksheet(remapped_workbook)
    for column_index, canonical in enumerate(canonical_columns, start=1):
        remapped_worksheet.cell(row=HEADER_ROW, column=column_index, value=canonical)

    for row_number in range(HEADER_ROW + 1, source_worksheet.max_row + 1):
        for column_index, canonical in enumerate(canonical_columns, start=1):
            source_column = source_column_by_canonical.get(canonical)
            value = (
                source_worksheet.cell(row=row_number, column=source_column).value
                if source_column is not None
                else None
            )
            remapped_worksheet.cell(row=row_number, column=column_index, value=value)

    return remapped_worksheet


def import_catalog_with_mapping(
    path: Path, column_mapping: Mapping[str, str]
) -> CatalogImportResult:
    """Importa un catálogo desde un archivo con cabeceras arbitrarias, vía ``column_mapping``."""
    workbook = open_workbook(path)
    worksheet = get_active_worksheet(workbook)
    remapped = build_remapped_worksheet(worksheet, column_mapping, PRODUCT_COLUMNS)
    return import_catalog_from_worksheet(remapped)


def import_packing_list_with_mapping(
    path: Path, column_mapping: Mapping[str, str], resolve_sku: Callable[[str], LoadUnit | None]
) -> PackingListImportResult:
    """Importa un Packing List desde un archivo con cabeceras arbitrarias vía ``column_mapping``."""
    workbook = open_workbook(path)
    worksheet = get_active_worksheet(workbook)
    remapped = build_remapped_worksheet(worksheet, column_mapping, PACKING_LIST_COLUMNS)
    return import_packing_list_from_worksheet(remapped, resolve_sku)


def import_loading_spaces_with_mapping(
    path: Path, column_mapping: Mapping[str, str]
) -> LoadingSpaceImportResult:
    """Importa Loading Spaces desde un archivo con cabeceras arbitrarias vía ``column_mapping``."""
    workbook = open_workbook(path)
    worksheet = get_active_worksheet(workbook)
    remapped = build_remapped_worksheet(worksheet, column_mapping, LOADING_SPACE_COLUMNS)
    return import_loading_spaces_from_worksheet(remapped)
