"""Utilidades comunes de apertura/escritura de libros `.xlsx`.

Centraliza dos cosas que todo importador/exportador necesita: abrir un
archivo detectando de forma clara si está corrupto o no es un `.xlsx`
válido, y escribirlo de forma atómica (mismo patrón que
`infrastructure/persistence/project_file_repository.py`, fase 7.0:
archivo temporal en el mismo directorio + `os.replace`, nunca escribir
directamente sobre la ruta final).
"""

from __future__ import annotations

import os
import tempfile
import zipfile
from collections.abc import Iterator, Sequence
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from cargo_optimizer.infrastructure.excel.exceptions import ExcelFileError, ExcelTemplateError

HEADER_ROW = 1
_FIRST_DATA_ROW = HEADER_ROW + 1


def open_workbook(path: Path) -> Workbook:
    """Abre un `.xlsx` existente, traduciendo cualquier fallo a `ExcelFileError`."""
    if not path.exists():
        raise ExcelFileError(f"El archivo '{path}' no existe.")
    try:
        return load_workbook(path, data_only=True, read_only=False)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ExcelFileError(
            f"No se pudo abrir '{path.name}': el archivo está corrupto o no es un .xlsx válido."
        ) from exc


def get_active_worksheet(workbook: Workbook) -> Worksheet:
    """Devuelve la hoja activa, garantizando que es una hoja de datos real.

    `Workbook.active` está tipado como `Worksheet | None` porque en
    teoría podría ser `None` o una hoja de gráficos — en la práctica,
    para cualquier `.xlsx` que abrimos nosotros mismos (nunca
    construido a mano por un tercero con solo gráficos), siempre hay al
    menos una hoja de datos real.
    """
    worksheet = workbook.active
    if not isinstance(worksheet, Worksheet):
        raise ExcelFileError("El archivo no contiene ninguna hoja de datos válida.")
    return worksheet


def require_headers(worksheet: Worksheet, expected_headers: Sequence[str]) -> None:
    """Valida que la primera fila coincida exactamente con las cabeceras esperadas.

    Comparación insensible a mayúsculas/minúsculas y a espacios sobrantes
    — dos plantillas que difieren solo en eso no deberían rechazarse.
    """
    actual = tuple(
        (
            ""
            if worksheet.cell(row=HEADER_ROW, column=i).value is None
            else str(worksheet.cell(row=HEADER_ROW, column=i).value).strip()
        )
        for i in range(1, len(expected_headers) + 1)
    )
    normalized_actual = tuple(value.casefold() for value in actual)
    normalized_expected = tuple(value.casefold() for value in expected_headers)
    if normalized_actual != normalized_expected:
        raise ExcelTemplateError(
            "El archivo no coincide con la plantilla esperada. "
            f"Cabeceras esperadas: {list(expected_headers)}. Cabeceras encontradas: {list(actual)}."
        )


def iter_data_rows(
    worksheet: Worksheet, *, column_count: int
) -> Iterator[tuple[int, tuple[Any, ...]]]:
    """Recorre las filas de datos (todo lo que sigue a la cabecera), saltando filas vacías.

    Devuelve ``(número_de_fila, valores)`` — el número de fila es el
    que vería un usuario en Excel (la cabecera es la fila 1), para que
    los mensajes de error sean directamente accionables.
    """
    for row_number in range(_FIRST_DATA_ROW, worksheet.max_row + 1):
        values = tuple(
            worksheet.cell(row=row_number, column=column).value
            for column in range(1, column_count + 1)
        )
        if all(value is None or str(value).strip() == "" for value in values):
            continue
        yield row_number, values


def save_workbook_atomic(workbook: Workbook, path: Path) -> None:
    """Escribe el libro en `path` de forma atómica (archivo temporal + `os.replace`)."""
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp_name = tempfile.mkstemp(
        prefix=f"{path.stem}.", suffix=".xlsx.tmp", dir=str(path.parent)
    )
    os.close(fd)
    tmp_path = Path(tmp_name)
    try:
        workbook.save(tmp_path)
        os.replace(tmp_path, path)
    except BaseException:
        tmp_path.unlink(missing_ok=True)
        raise
