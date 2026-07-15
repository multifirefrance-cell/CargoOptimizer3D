"""Exportar un `PackingResult` completo a un `.xlsx` profesional.

Cinco hojas en un único libro: Resumen (incluye volumen, peso,
porcentajes, tiempo y algoritmo), Productos cargados, Productos no
cargados, Warnings y Datos del espacio. Sin macros, sin VBA.

``load_units_by_id`` se recibe junto al `PackingResult` para resolver
SKU/nombre de cada `Placement`/`UnpackedUnit` — mismo patrón ya
establecido para `Packing3DViewer.display_result` (fase 6.1, ver
`docs/ThreeDViewer.md`) y `UnpackedUnitTableModel` (fase 5.1): evita
ampliar `PackingResult`/`Placement` con datos de `LoadUnit` que no les
corresponden.
"""

from __future__ import annotations

from collections.abc import Mapping
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from uuid import UUID

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from cargo_optimizer.domain.load_unit import LoadUnit
from cargo_optimizer.domain.packing_result import PackingResult
from cargo_optimizer.infrastructure.excel.loading_space_rows import (
    CATEGORY_LABELS,
    DOOR_POSITION_LABELS,
)
from cargo_optimizer.infrastructure.excel.product_rows import ORIENTATION_LABELS
from cargo_optimizer.infrastructure.excel.styles import (
    LABEL_FONT,
    TITLE_FONT,
    apply_borders_to_data_rows,
    apply_table,
    autofit_columns,
    write_header_row,
)
from cargo_optimizer.infrastructure.excel.workbook_utils import (
    HEADER_ROW,
    get_active_worksheet,
    save_workbook_atomic,
)

SHEET_SUMMARY = "Resumen"
SHEET_PACKED = "Productos cargados"
SHEET_UNPACKED = "Productos no cargados"
SHEET_WARNINGS = "Warnings"
SHEET_SPACE = "Datos del espacio"

_PACKED_COLUMNS = (
    "SKU",
    "Nombre",
    "Instancia",
    "Posicion X (cm)",
    "Posicion Y (cm)",
    "Posicion Z (cm)",
    "Orientacion",
    "Largo colocado (cm)",
    "Ancho colocado (cm)",
    "Alto colocado (cm)",
    "Orden de colocacion",
)
_UNPACKED_COLUMNS = ("SKU", "Nombre", "Instancia", "Motivo", "Detalle")
_WARNING_COLUMNS = ("Aviso",)


def _unit_label(load_units_by_id: Mapping[UUID, LoadUnit], unit_id: UUID) -> tuple[str, str]:
    unit = load_units_by_id.get(unit_id)
    if unit is None:
        return (str(unit_id), "(producto desconocido)")
    return (unit.sku, unit.name)


def _write_label_value_row(worksheet: Worksheet, row: int, label: str, value: Any) -> int:
    label_cell = worksheet.cell(row=row, column=1, value=label)
    label_cell.font = LABEL_FONT
    worksheet.cell(row=row, column=2, value=value)
    return row + 1


def _build_summary_sheet(
    worksheet: Worksheet, result: PackingResult, *, application_version: str
) -> None:
    worksheet.cell(row=1, column=1, value="Resumen de la optimizacion").font = TITLE_FONT
    row = 3
    weight_utilization = result.weight_utilization_percent
    rows: tuple[tuple[str, object], ...] = (
        ("Algoritmo utilizado", result.algorithm_name),
        ("Tiempo de ejecucion (s)", round(result.execution_time_seconds, 3)),
        ("Unidades solicitadas", result.requested_count),
        ("Unidades cargadas", result.packed_count),
        ("Unidades no cargadas", result.unpacked_count),
        ("Porcentaje completado (%)", round(result.packing_completion_percent, 2)),
        ("Volumen utilizado (cm3)", round(result.used_volume_cm3, 2)),
        ("Volumen utilizado (m3)", round(result.used_volume_cm3 / 1_000_000.0, 4)),
        ("Porcentaje de volumen utilizado (%)", round(result.volume_utilization_percent, 2)),
        ("Peso utilizado (kg)", round(result.used_weight_kg, 2)),
        (
            "Porcentaje de peso utilizado (%)",
            "Sin limite declarado" if weight_utilization is None else round(weight_utilization, 2),
        ),
        ("Numero de avisos", len(result.warnings)),
        ("Generado con CargoOptimizer3D version", application_version),
        ("Fecha de generacion", datetime.now(UTC).strftime("%Y-%m-%d %H:%M:%S UTC")),
    )
    for label, value in rows:
        row = _write_label_value_row(worksheet, row, label, value)

    worksheet.column_dimensions["A"].width = 36
    worksheet.column_dimensions["B"].width = 30


def _build_packed_sheet(
    worksheet: Worksheet, result: PackingResult, load_units_by_id: Mapping[UUID, LoadUnit]
) -> None:
    write_header_row(worksheet, _PACKED_COLUMNS, row=HEADER_ROW)
    last_row = HEADER_ROW
    for placement in sorted(result.placements, key=lambda p: p.sequence_number):
        sku, name = _unit_label(load_units_by_id, placement.load_unit_id)
        last_row += 1
        worksheet.append(
            (
                sku,
                name,
                placement.instance_number,
                round(placement.x_cm, 2),
                round(placement.y_cm, 2),
                round(placement.z_cm, 2),
                ORIENTATION_LABELS[placement.orientation.code],
                round(placement.length_cm, 2),
                round(placement.width_cm, 2),
                round(placement.height_cm, 2),
                placement.sequence_number,
            )
        )
    _finish_table_sheet(worksheet, "TablaProductosCargados", len(_PACKED_COLUMNS), last_row)


def _build_unpacked_sheet(
    worksheet: Worksheet, result: PackingResult, load_units_by_id: Mapping[UUID, LoadUnit]
) -> None:
    write_header_row(worksheet, _UNPACKED_COLUMNS, row=HEADER_ROW)
    last_row = HEADER_ROW
    for unpacked in result.unpacked_units:
        sku, name = _unit_label(load_units_by_id, unpacked.load_unit_id)
        last_row += 1
        worksheet.append(
            (sku, name, unpacked.instance_number, unpacked.reason_code, unpacked.reason_message)
        )
    _finish_table_sheet(worksheet, "TablaProductosNoCargados", len(_UNPACKED_COLUMNS), last_row)


def _build_warnings_sheet(worksheet: Worksheet, result: PackingResult) -> None:
    write_header_row(worksheet, _WARNING_COLUMNS, row=HEADER_ROW)
    last_row = HEADER_ROW
    for warning in result.warnings:
        last_row += 1
        worksheet.append((warning,))
    _finish_table_sheet(worksheet, "TablaWarnings", len(_WARNING_COLUMNS), last_row)
    worksheet.column_dimensions["A"].width = 80


def _build_space_sheet(worksheet: Worksheet, result: PackingResult) -> None:
    space = result.loading_space
    worksheet.cell(row=1, column=1, value="Datos del espacio de carga").font = TITLE_FONT
    row = 3
    rows: tuple[tuple[str, object], ...] = (
        ("Nombre", space.name),
        ("Categoria", CATEGORY_LABELS[space.category]),
        ("Largo (cm)", space.internal_dimensions.length_cm),
        ("Ancho (cm)", space.internal_dimensions.width_cm),
        ("Alto (cm)", space.internal_dimensions.height_cm),
        ("Volumen (cm3)", round(space.capacity_volume_cm3, 2)),
        ("Volumen (m3)", round(space.capacity_volume_m3, 4)),
        (
            "Peso maximo (kg)",
            "Sin limite declarado" if space.max_weight_kg is None else space.max_weight_kg,
        ),
        ("Posicion de puerta", DOOR_POSITION_LABELS[space.door_position]),
        ("Notas", space.notes),
    )
    for label, value in rows:
        row = _write_label_value_row(worksheet, row, label, value)

    worksheet.column_dimensions["A"].width = 24
    worksheet.column_dimensions["B"].width = 40


def _finish_table_sheet(
    worksheet: Worksheet, table_name: str, column_count: int, last_row: int
) -> None:
    apply_borders_to_data_rows(
        worksheet, first_row=HEADER_ROW, last_row=last_row, column_count=column_count
    )
    apply_table(
        worksheet,
        table_name=table_name,
        first_row=HEADER_ROW + 1,
        last_row=last_row,
        column_count=column_count,
    )
    autofit_columns(worksheet, column_count=column_count)


def export_packing_result(
    result: PackingResult,
    load_units_by_id: Mapping[UUID, LoadUnit],
    path: Path,
    *,
    application_version: str,
) -> None:
    """Escribe `result` a `path` con el formato de `OptimizationResultTemplate.xlsx`."""
    workbook = Workbook()

    summary_sheet = get_active_worksheet(workbook)
    summary_sheet.title = SHEET_SUMMARY
    _build_summary_sheet(summary_sheet, result, application_version=application_version)

    _build_packed_sheet(workbook.create_sheet(SHEET_PACKED), result, load_units_by_id)
    _build_unpacked_sheet(workbook.create_sheet(SHEET_UNPACKED), result, load_units_by_id)
    _build_warnings_sheet(workbook.create_sheet(SHEET_WARNINGS), result)
    _build_space_sheet(workbook.create_sheet(SHEET_SPACE), result)

    save_workbook_atomic(workbook, path)
